import os

os.environ["DATABASE_URL"] = "sqlite://"

from openpyxl import Workbook
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import inspect as sqlalchemy_inspect

from app.database import Base, SessionLocal, engine, migrate_legacy_schema
from app.main import app
from app.models import (
    AdministrativeBoundary,
    AnalysisRun,
    CensusIndicator,
    City,
    DistrictCrosswalk,
    GeographicPlace,
    MarketFactor,
    RetailLocation,
)
from app.services.census_import import import_census_workbook


@pytest.fixture
def client():
    with TestClient(app) as test_client:
        yield test_client


def test_health_and_city_list_use_api_prefix(client):
    assert client.get("/api/health").json() == {"status": "healthy"}
    city_response = client.get("/api/cities")
    assert city_response.json() == {"cities": []}
    assert city_response.history == []


def test_database_models_register_once_and_initialize(client):
    assert set(Base.metadata.tables) == {
        "cities",
        "market_factors",
        "analysis_runs",
        "retail_locations",
        "census_indicators",
        "geographic_places",
        "administrative_boundaries",
        "dataset_registry",
        "district_crosswalks",
    }
    assert (
        sqlalchemy_inspect(City).relationships["factors"].mapper.class_
        is MarketFactor
    )
    assert set(sqlalchemy_inspect(engine).get_table_names()) >= {
        "cities",
        "market_factors",
        "analysis_runs",
        "retail_locations",
        "census_indicators",
        "geographic_places",
        "administrative_boundaries",
        "dataset_registry",
        "district_crosswalks",
    }
    assert AnalysisRun.__table__.metadata is Base.metadata


def test_location_routes_are_registered_and_disclose_missing_data(client):
    response = client.get("/api/locations")

    assert response.status_code == 200
    assert response.json() == {
        "count": 0,
        "locations": [],
        "data_status": "unavailable",
    }

    filtered = client.get(
        "/api/locations",
        params={"city": "Pune", "category": "Grocery"},
    )
    assert filtered.status_code == 200
    assert filtered.json()["locations"] == []

    missing_location = client.get("/api/locations/999999")
    assert missing_location.status_code == 404


def test_geography_places_endpoint_discloses_missing_import(client):
    with SessionLocal() as db:
        db.query(GeographicPlace).delete()
        db.commit()

    response = client.get("/api/geography/places")
    assert response.status_code == 200
    assert response.json()["data_status"] == "unavailable"
    assert response.json()["count"] == 0
    assert response.json()["total"] == 0
    assert response.json()["snapshot_date"] is None


def test_geography_boundaries_endpoint_discloses_missing_import(client):
    with SessionLocal() as db:
        db.query(DistrictCrosswalk).delete()
        db.query(AdministrativeBoundary).delete()
        db.commit()

    response = client.get("/api/geography/boundaries")

    assert response.status_code == 200
    assert response.json()["data_status"] == "unavailable"
    assert response.json()["count"] == 0
    assert response.json()["total"] == 0
    assert response.json()["reported_source_count"] == 736
    assert "official district codes" in response.json()["coverage_note"]
    assert client.get("/api/geography/boundaries/geojson").status_code == 404


def test_location_routes_are_in_openapi_schema(client):
    paths = client.get("/openapi.json").json()["paths"]

    assert "/api/locations" in paths
    assert "/api/locations/{location_id}" in paths
    assert "/api/locations/{location_id}/nearby" in paths
    assert "/api/cities/factors" in paths
    assert "/api/census/indicators" in paths
    assert "/api/geography/places" in paths
    assert "/api/geography/boundaries" in paths
    assert "/api/datasets" in paths
    assert "/api/geography/audit" in paths
    assert "/api/market-scores/evaluate" in paths


def test_comparison_does_not_use_census_population_as_a_retail_score(client):
    with SessionLocal() as db:
        db.add(
            CensusIndicator(
                geography_code="01|000|00000||500400000|0|Srinagar UA",
                geography_name="Srinagar UA",
                state_code="01",
                ua_code="500400000",
                census_level=0,
                indicator_code="TOT_P",
                indicator_name="Population - Total",
                value=1_264_202,
                unit="person",
                reference_year=2011,
                source_name="Census of India",
                source_url="https://example.test/census",
                download_url="https://example.test/census.xlsx",
                license_status="not_specified",
            )
        )
        db.commit()
        try:
            response = client.post(
                "/api/comparisons",
                json={
                    "category": "Grocery",
                    "cities": ["Srinagar UA", "Pune"],
                },
            )
            assert response.status_code == 200
            assert response.json()["data_status"] == "no_complete_scores"
            assert all(
                result["total_score"] is None
                for result in response.json()["results"]
            )
        finally:
            db.query(CensusIndicator).delete()
            db.commit()


def test_census_import_is_validated_idempotent_and_retrievable(client, tmp_path):
    workbook_path = tmp_path / "sample-census.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(
        [
            "ST Code",
            "DT Code",
            "SubDT Code",
            "Town Code",
            "UA Code",
            "UA Name",
            "Level",
            "TOT_P",
        ]
    )
    records = [
        ("01", "001", "00001", None, "UA001", "Alpha UA", 0, 100),
        ("01", "001", "00001", "000001", "UA002", "Beta town", 1, 50),
        ("01", "001", "00001", "000002", "UA003", "Missing total", None, None),
        ("01", "001", "00001", None, "UA001", "Alpha UA", 0, 100),
        ("01", "001", "00001", "000003", "UA004", "Invalid total", 1, -1),
    ]
    for record in records:
        sheet.append(record)
    workbook.save(workbook_path)

    db = SessionLocal()
    try:
        db.query(CensusIndicator).delete()
        db.commit()

        first_import = import_census_workbook(db, workbook_path)
        assert first_import["inserted"] == 3
        assert first_import["updated"] == 0
        assert first_import["skipped"] == 2
        assert first_import["missing_values"] == 1
        assert len(first_import["validation_errors"]) == 2

        second_import = import_census_workbook(db, workbook_path)
        assert second_import["inserted"] == 0
        assert second_import["updated"] == 3

        response = client.get(
            "/api/census/indicators",
            params={"search": "Missing", "state_code": "01"},
        )
        assert response.status_code == 200
        body = response.json()
        assert body["total"] == 1
        assert body["dataset_total"] == 3
        assert body["reference_year"] == 2011
        assert body["license_status"] == "not_specified_in_catalog"
        assert body["data_status"] == "available"
        assert body["indicators"][0]["value"] is None
        assert body["indicators"][0]["geography_name"] == "Missing total"

        level_filter = client.get(
            "/api/census/indicators",
            params={"census_level": 0},
        )
        assert level_filter.json()["total"] == 1

        unreported_filter = client.get(
            "/api/census/indicators",
            params={"level_unreported": "true"},
        )
        assert unreported_filter.json()["total"] == 1
    finally:
        db.query(CensusIndicator).delete()
        db.commit()
        db.close()


def test_census_import_rejects_workbooks_missing_required_columns(tmp_path):
    workbook_path = tmp_path / "incomplete-census.xlsx"
    workbook = Workbook()
    workbook.active.title = "Data"
    workbook.active.append(["UA Name", "TOT_P"])
    workbook.save(workbook_path)

    db = SessionLocal()
    try:
        with pytest.raises(ValueError, match="missing required columns"):
            import_census_workbook(db, workbook_path)
    finally:
        db.close()


def test_legacy_location_schema_migration_preserves_rows(tmp_path):
    from sqlalchemy import create_engine, text

    legacy_engine = create_engine(f"sqlite:///{tmp_path / 'legacy.db'}")
    try:
        with legacy_engine.begin() as connection:
            connection.execute(
                text(
                    "CREATE TABLE retail_locations ("
                    "id INTEGER PRIMARY KEY, name VARCHAR(160) NOT NULL)"
                )
            )
            connection.execute(
                text("INSERT INTO retail_locations (id, name) VALUES (1, 'Old site')")
            )

        migrate_legacy_schema(legacy_engine)

        with legacy_engine.connect() as connection:
            columns = {
                column["name"]
                for column in sqlalchemy_inspect(connection).get_columns(
                    "retail_locations"
                )
            }
            row = connection.execute(
                text(
                    "SELECT name, location_type FROM retail_locations WHERE id = 1"
                )
            ).one()
        assert "location_type" in columns
        assert row == ("Old site", "candidate")
    finally:
        legacy_engine.dispose()


@pytest.mark.parametrize(
    "origin",
    [
        "http://localhost:5173",
        "http://127.0.0.1:5173",
        "http://localhost:4173",
        "http://127.0.0.1:4173",
    ],
)
def test_frontend_origin_can_call_the_api(client, origin):
    response = client.options(
        "/api/comparisons",
        headers={
            "Origin": origin,
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "content-type",
        },
    )

    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == origin


def test_comparison_returns_missing_data_without_inventing_scores(client):
    response = client.post(
        "/api/comparisons",
        json={"category": "Grocery", "cities": ["Pune", "Mumbai"]},
    )

    assert response.status_code == 200
    assert response.history == []
    body = response.json()
    assert body["category"] == "Grocery"
    assert body["data_status"] == "no_complete_scores"
    assert body["analysis_id"] > 0
    assert {result["city_name"] for result in body["results"]} == {
        "Pune",
        "Mumbai",
    }
    assert all(result["total_score"] is None for result in body["results"])
    assert all(
        result["source_status"] == "no_matching_data"
        for result in body["results"]
    )


def test_comparison_scores_and_persists_matching_factor_data(client, monkeypatch):
    from app.routers import comparisons

    monkeypatch.setattr(
        comparisons,
        "get_cities_for_category",
        lambda _category, db=None: [
            {
                "city": "Pune",
                "category": "Grocery",
                "demand": 80,
                "competition_opportunity": 70,
                "accessibility": 90,
                "cost_viability": 75,
                "category_fit": 85,
                "source": "Verified local study",
                "source_status": "verified",
            }
        ],
    )

    response = client.post(
        "/api/comparisons",
        json={"category": "Grocery", "cities": ["Pune", "Mumbai"]},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["data_status"] == "available"
    assert body["analysis_id"] > 0

    pune, mumbai = body["results"]
    assert pune["city_name"] == "Pune"
    assert pune["total_score"] == 80.25
    assert pune["source_status"] == "verified"
    assert pune["source"] == "Verified local study"
    assert pune["score_status"] == "complete"
    assert pune["rank"] == 1
    assert mumbai["city_name"] == "Mumbai"
    assert mumbai["total_score"] is None
    assert mumbai["source_status"] == "no_matching_data"
    assert mumbai["source"] is None
    assert mumbai["rank"] is None


def test_comparison_withholds_complete_but_unverified_factor_scores(
    client, monkeypatch
):
    from app.routers import comparisons

    monkeypatch.setattr(
        comparisons,
        "get_cities_for_category",
        lambda _category, db=None: [
            {
                "city": "Pune",
                "category": "Grocery",
                "demand": 80,
                "competition_opportunity": 70,
                "accessibility": 90,
                "cost_viability": 75,
                "category_fit": 85,
                "source": "User-entered estimates",
                "source_status": "user_provided",
            }
        ],
    )

    response = client.post(
        "/api/comparisons",
        json={"category": "Grocery", "cities": ["Pune", "Mumbai"]},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["data_status"] == "no_complete_scores"
    pune = next(item for item in body["results"] if item["city_name"] == "Pune")
    assert pune["total_score"] is None
    assert pune["score_status"] == "incomplete"
    assert "verified_source" in pune["missing_factors"]
    assert pune["source_status"] == "user_provided"
    assert pune["rank"] is None


def test_comparison_rejects_duplicate_city_names(client):
    response = client.post(
        "/api/comparisons",
        json={"category": "Grocery", "cities": ["Pune", " pune "]},
    )

    assert response.status_code == 422


def test_scenario_and_csv_endpoints_accept_frontend_contract(client):
    scenario_response = client.post(
        "/api/scenarios/analyze",
        json={
            "category": "Grocery",
            "factors": {
                "demand": 50,
                "competition_opportunity": 70,
                "accessibility": 80,
                "cost_viability": 60,
                "category_fit": 90,
            },
            "demand_change": 10,
            "cost_change": 10,
        },
    )
    assert scenario_response.status_code == 200
    assert scenario_response.json()["adjusted_factors"]["cost_viability"] == 50

    export_response = client.post(
        "/api/exports/csv",
        json={
            "results": [
                {
                    "city_name": "Pune",
                    "category": "Grocery",
                    "total_score": None,
                    "missing_factors": ["demand", "cost_viability"],
                    "source": "Reviewed report",
                    "notes": "Evidence note",
                }
            ]
        },
    )
    assert export_response.status_code == 200
    assert "Pune,Grocery,,," in export_response.text
    assert "demand; cost_viability" in export_response.text
    assert "Reviewed report" in export_response.text
    assert "Evidence note" in export_response.text


def test_market_factor_crud_drives_city_list_and_comparison(client):
    payload = {
        "city": "Pune",
        "category": "Grocery",
        "demand": 80,
        "competition_opportunity": 70,
        "accessibility": 90,
        "cost_viability": 75,
        "category_fit": 85,
        "source": "Verified local study",
        "source_status": "verified",
        "notes": "Evidence reviewed",
    }

    saved = client.post("/api/cities/factors", json=payload)
    assert saved.status_code == 200
    assert saved.json()["source_status"] == "verified"
    assert saved.json()["notes"] == "Evidence reviewed"
    assert client.get("/api/cities").json() == {"cities": ["Pune"]}
    assert client.get("/api/cities/factors").json()[0]["city"] == "Pune"

    comparison = client.post(
        "/api/comparisons",
        json={"category": "Grocery", "cities": ["Pune", "Mumbai"]},
    )
    results = {item["city_name"]: item for item in comparison.json()["results"]}
    assert results["Pune"]["total_score"] == 80.25
    assert results["Pune"]["source_status"] == "verified"
    assert results["Pune"]["source"] == "Verified local study"
    assert results["Pune"]["notes"] == "Evidence reviewed"
    assert results["Mumbai"]["total_score"] is None

    deleted = client.delete(f"/api/cities/factors/{saved.json()['id']}")
    assert deleted.status_code == 204
    assert client.get("/api/cities/factors").json() == []
    assert client.get("/api/cities").json() == {"cities": []}


def test_market_factor_cannot_be_marked_verified_without_source(client):
    response = client.post(
        "/api/cities/factors",
        json={
            "city": "Pune",
            "category": "Grocery",
            "demand": 80,
            "source_status": "verified",
        },
    )

    assert response.status_code == 422


def test_location_crud_filters_and_geographic_distance(client):
    first = client.post(
        "/api/locations",
        json={
            "name": "Candidate A",
            "city": "Pune",
            "category": "Grocery",
            "latitude": 18.5204,
            "longitude": 73.8567,
            "source_status": "user_provided",
        },
    )
    second = client.post(
        "/api/locations",
        json={
            "name": "Candidate B",
            "city": "Pune",
            "location_type": "competitor",
            "category": "Grocery",
            "latitude": 18.53,
            "longitude": 73.86,
        },
    )
    assert first.status_code == second.status_code == 201

    location_id = first.json()["id"]
    filtered = client.get(
        "/api/locations",
        params={"city": "pune", "category": "Grocery"},
    )
    assert filtered.json()["count"] == 2

    nearby = client.get(f"/api/locations/{location_id}/nearby")
    assert nearby.status_code == 200
    assert nearby.json()["count"] == 1
    assert nearby.json()["nearby_type"] == "competitor"
    assert nearby.json()["locations"][0]["name"] == "Candidate B"
    assert nearby.json()["locations"][0]["distance_km"] > 0

    updated = client.put(
        f"/api/locations/{location_id}",
        json={"name": "Updated candidate"},
    )
    assert updated.json()["name"] == "Updated candidate"
    null_update = client.put(
        f"/api/locations/{location_id}",
        json={"latitude": None},
    )
    assert null_update.status_code == 422
    assert client.delete(f"/api/locations/{location_id}").status_code == 204
    assert client.get(f"/api/locations/{location_id}").status_code == 404


def test_market_and_location_inputs_validate_ranges_and_required_values(client):
    assert client.post(
        "/api/cities/factors",
        json={"city": "Pune", "category": "Grocery"},
    ).status_code == 422
    assert client.post(
        "/api/locations",
        json={
            "name": "Invalid",
            "city": "Pune",
            "category": "Grocery",
            "latitude": 91,
            "longitude": 74,
        },
    ).status_code == 422


def test_scenario_rejects_boolean_factor_values(client):
    response = client.post(
        "/api/scenarios/analyze",
        json={
            "category": "Grocery",
            "factors": {
                "demand": True,
                "competition_opportunity": 70,
                "accessibility": 80,
                "cost_viability": 60,
                "category_fit": 90,
            },
        },
    )

    assert response.status_code == 422
