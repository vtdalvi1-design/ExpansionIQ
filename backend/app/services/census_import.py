import argparse
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.database import Base, SessionLocal, engine
from app.models import CensusIndicator
from app.routers.census import (
    DOWNLOAD_URL,
    LICENSE_STATUS,
    REFERENCE_YEAR,
    SOURCE_NAME,
    SOURCE_URL,
)

REQUIRED_COLUMNS = {
    "ST Code",
    "DT Code",
    "SubDT Code",
    "Town Code",
    "UA Code",
    "UA Name",
    "Level",
    "TOT_P",
}


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _population(value: Any, row_number: int) -> int | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"Row {row_number}: TOT_P must be a non-negative integer.")
    if value < 0 or int(value) != value:
        raise ValueError(f"Row {row_number}: TOT_P must be a non-negative integer.")
    return int(value)


def _records(rows: Iterable[tuple[Any, ...]]) -> Iterable[tuple[int, dict[str, Any]]]:
    try:
        header = next(iter(rows))
    except StopIteration as exc:
        raise ValueError("The workbook Data sheet is empty.") from exc

    header_positions = {
        str(column).strip(): index
        for index, column in enumerate(header)
        if column is not None
    }
    missing_columns = REQUIRED_COLUMNS - header_positions.keys()
    if missing_columns:
        raise ValueError(
            "The workbook is missing required columns: "
            + ", ".join(sorted(missing_columns))
        )

    for row_number, row in enumerate(rows, start=2):
        yield row_number, {
            column: row[index] if index < len(row) else None
            for column, index in header_positions.items()
            if column in REQUIRED_COLUMNS
        }


def import_census_workbook(
    db: Session,
    workbook_path: Path,
) -> dict[str, Any]:
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Census workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Data" not in workbook.sheetnames:
        raise ValueError("The Census workbook does not contain a Data sheet.")

    inserted = 0
    updated = 0
    skipped = 0
    missing_values = 0
    validation_errors: list[str] = []
    seen_codes: set[str] = set()

    try:
        rows = workbook["Data"].iter_rows(values_only=True)
        for row_number, record in _records(rows):
            if not any(value is not None for value in record.values()):
                continue

            state_code = _text(record.get("ST Code"))
            ua_code = _text(record.get("UA Code"))
            geography_name = _text(record.get("UA Name"))
            if not state_code or not ua_code or not geography_name:
                validation_errors.append(
                    f"Row {row_number}: state code, UA code, and geography name are required."
                )
                skipped += 1
                continue

            census_level = record.get("Level")
            if census_level is not None and (
                isinstance(census_level, bool)
                or not isinstance(census_level, int)
                or not 0 <= census_level <= 3
            ):
                validation_errors.append(
                    f"Row {row_number}: Level must be an integer from 0 to 3 or blank."
                )
                skipped += 1
                continue

            try:
                population = _population(record.get("TOT_P"), row_number)
            except ValueError as error:
                validation_errors.append(str(error))
                skipped += 1
                continue

            geographic_parts = (
                state_code,
                _text(record.get("DT Code")) or "",
                _text(record.get("SubDT Code")) or "",
                _text(record.get("Town Code")) or "",
                ua_code,
                "" if census_level is None else str(census_level),
                geography_name,
            )
            geography_code = "|".join(geographic_parts)
            if geography_code in seen_codes:
                validation_errors.append(
                    f"Row {row_number}: duplicate geography record {geography_code}."
                )
                skipped += 1
                continue
            seen_codes.add(geography_code)

            existing = (
                db.query(CensusIndicator)
                .filter(
                    CensusIndicator.geography_code == geography_code,
                    CensusIndicator.indicator_code == "TOT_P",
                    CensusIndicator.reference_year == REFERENCE_YEAR,
                    CensusIndicator.source_url == SOURCE_URL,
                )
                .first()
            )
            if existing is None:
                existing = CensusIndicator(
                    geography_code=geography_code,
                    indicator_code="TOT_P",
                    reference_year=REFERENCE_YEAR,
                    source_url=SOURCE_URL,
                )
                db.add(existing)
                inserted += 1
            else:
                updated += 1

            existing.geography_name = geography_name
            existing.state_code = state_code
            existing.district_code = _text(record.get("DT Code"))
            existing.subdistrict_code = _text(record.get("SubDT Code"))
            existing.town_code = _text(record.get("Town Code"))
            existing.ua_code = ua_code
            existing.census_level = census_level
            existing.indicator_name = "Population - Total"
            existing.value = population
            existing.unit = "person"
            existing.source_name = SOURCE_NAME
            existing.download_url = DOWNLOAD_URL
            existing.license_status = LICENSE_STATUS
            if population is None:
                missing_values += 1

        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        workbook.close()

    return {
        "inserted": inserted,
        "updated": updated,
        "skipped": skipped,
        "missing_values": missing_values,
        "validation_errors": validation_errors,
    }


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Import official 2011 Census urban population records."
    )
    parser.add_argument(
        "workbook",
        type=Path,
        help="Path to the downloaded PCA11-UA-0000.xlsx workbook.",
    )
    args = parser.parse_args()

    Base.metadata.create_all(bind=engine)
    with SessionLocal() as db:
        summary = import_census_workbook(db, args.workbook)
    print(
        "Census import complete: "
        f"{summary['inserted']} inserted, {summary['updated']} updated, "
        f"{summary['skipped']} skipped, "
        f"{summary['missing_values']} population values missing."
    )
    for error in summary["validation_errors"]:
        print(f"Validation issue: {error}")


if __name__ == "__main__":
    main()
